"""Ensambla las escenas en un archivo .pptx.

Un único registro de identidad recorre toda la compilación, de modo que
un objeto presente en varias escenas conserve su identificador OOXML en
todas las slides. Esa estabilidad es lo que hace posible el Morph.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

from pptx_compiler.compiler.differ import diff
from pptx_compiler.ir.identity import IdentityRegistry
from pptx_compiler.ir.scene import Scene
from pptx_compiler.render.projection import SLIDE_H_EMU, SLIDE_W_EMU
from pptx_compiler.render.shapes import draw_scene
from pptx_compiler.render.transitions import DEFAULT_DURATION_MS, apply_transition

BLANK_LAYOUT = 6


def build(
    scenes: list[Scene],
    world_w: float,
    output: Path,
    title: str = "",
    transition_ms: int = DEFAULT_DURATION_MS,
) -> Path:
    presentation = Presentation()
    presentation.slide_width = Emu(SLIDE_W_EMU)
    presentation.slide_height = Emu(SLIDE_H_EMU)
    presentation.core_properties.title = title

    identity = IdentityRegistry()
    layout = presentation.slide_layouts[BLANK_LAYOUT]

    slides = []
    for scene in scenes:
        slide = presentation.slides.add_slide(layout)
        draw_scene(slide, scene, identity, world_w)
        slides.append(slide)

    _apply_transitions(slides, scenes, transition_ms)

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))
    return output


def _apply_transitions(
    slides: list, scenes: list[Scene], transition_ms: int
) -> None:
    """La transición se aplica a la slide de destino, no a la de origen."""
    for index in range(1, len(scenes)):
        result = diff(scenes[index - 1], scenes[index])
        apply_transition(slides[index], result.kind, duration_ms=transition_ms)
